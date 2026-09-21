import os
import re
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from typing import List, Optional
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import latex2mathml.converter
import lxml.etree as LET

app = FastAPI(title="FGEI Multi-Format Exam Generator")

UPLOAD_FOLDER = "static/uploads"
os.makedirs(
    UPLOAD_FOLDER,
    exist_ok=True
)

@app.post("/upload-image")
async def upload_image(
    file: UploadFile = File(...)
):

    allowed = [
        ".png",
        ".jpg",
        ".jpeg"
    ]

    ext = os.path.splitext(file.filename)[1].lower()


    if ext not in allowed:
        raise HTTPException(
            status_code=400,
            detail="Only PNG/JPG images allowed"
        )


    file_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )


    with open(file_path, "wb") as buffer:

        buffer.write(
            await file.read()
        )


    return {
        "image": file_path.replace("\\","/")
    }

if os.path.exists("static"):

    app.mount("/static", StaticFiles(directory="static"), name="static")

class MCQQuestion(BaseModel):

    text: str

    is_urdu: Optional[bool] = False

    options: List[str]

    image: Optional[str] = None

class QuestionItem(BaseModel):

    main_question: str

    or_question: Optional[str] = None

    is_urdu: Optional[bool] = False

    marks: Optional[int] = 3

    image: Optional[str] = None

class ExamData(BaseModel):

    school_name: str
    exam_title: str
    subject: str
    grade: str
    time_allowed: str
    total_marks: int

    # Template control
    is_board_pattern: bool = False
    is_urdu_paper: bool = False

    # Future optional manual override
    paper_template: Optional[str] = None

    mcqs: List[MCQQuestion] = Field(default_factory=list)

    short_questions: List[QuestionItem] = Field(default_factory=list)

    long_questions: List[QuestionItem] = Field(default_factory=list)

    project_questions: List[QuestionItem] = Field(default_factory=list)

def detect_template(data: ExamData):

    if data.is_urdu_paper:
        return "URDU"

    if data.paper_template:
        return data.paper_template.upper()

    grade = data.grade.lower()

    if "9" in grade or "10" in grade:
        return "SSC"

    return "PRIMARY"    

# Standard MS Word OMML Transformation Sheet

with open("MML2OMML.XSL", "r", encoding="utf-8") as f:
    MML2OMML_XSLT = f.read()

def latex_to_omml(latex_str: str):

    try:

        clean_latex = latex_str.strip()


        # Convert LaTeX to MathML
        mathml_str = latex2mathml.converter.convert(
            clean_latex
        )


        # Parse MathML
        mathml_tree = LET.fromstring(
            mathml_str.encode("utf-8")
        )


        # Load Microsoft's MathML → OMML converter
        xslt_root = LET.XML(
            MML2OMML_XSLT.encode("utf-8")
        )


        transform = LET.XSLT(
            xslt_root
        )


        # Generate real Word equation XML
        omml_tree = transform(
            mathml_tree
        )

        xml = LET.tostring(
    omml_tree,
    encoding="unicode"
)

        return parse_xml(
            xml.encode("utf-8")
    
)


    except Exception as e:

        print(
            "Equation conversion error:",
            e
        )

        return None

def add_text_with_math(paragraph, full_text: str):

    if not full_text:
        return


    pattern = r"\[MATH\](.*?)\[/MATH\]"


    parts = re.split(
        pattern,
        full_text,
        flags=re.DOTALL
    )


    for index, part in enumerate(parts):

        # Odd indexes are math sections
        if index % 2 == 1:

            omml = latex_to_omml(part)

            if omml is not None:

                paragraph._element.append(
                    omml
                )

        else:

            if part:

                paragraph.add_run(part)

def set_paragraph_rtl(paragraph, font_name="Jameel Noori Nastaliq", font_size=11, align_right=True):

    pPr = paragraph._element.get_or_add_pPr()

    bidi = OxmlElement('w:bidi')

    pPr.append(bidi)

    if align_right:

        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for run in paragraph.runs:

        run.font.name = font_name

        run.font.size = Pt(font_size)

def build_fgei_paper(data: ExamData, output_filename="fgei_paper.docx"):
    template = detect_template(data)


    if template == "SSC":
        return build_ssc_paper(data, output_filename)


    elif template == "PRIMARY":
        return build_primary_paper(data, output_filename)


    elif template == "URDU":
        return build_urdu_paper(data, output_filename)


    doc = Document()

    for section in doc.sections:

        section.top_margin = Inches(0.5)

        section.bottom_margin = Inches(0.5)

        section.left_margin = Inches(0.5)

        section.right_margin = Inches(0.5)

    # 1. Header & Logo Formatting

    logo_path = os.path.join("static", "logo.png")

    if os.path.exists(logo_path):

        p_logo = doc.add_paragraph()

        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_logo.paragraph_format.space_after = Pt(2)

        p_logo.add_run().add_picture(logo_path, width=Inches(0.85))

    p_head = doc.add_paragraph()

    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_head.paragraph_format.space_after = Pt(8)

    r_sch = p_head.add_run(f"{data.school_name}\n")

    r_sch.bold = True

    r_sch.font.size = Pt(13)

    r_ttl = p_head.add_run(f"{data.exam_title}\n")

    r_ttl.bold = True

    r_ttl.font.size = Pt(11)

    p_head.add_run(f"Subject: {data.subject}  |  Class: {data.grade}  |  Time: {data.time_allowed}  |  Total Marks: {data.total_marks}")

    if data.is_urdu_paper:

        set_paragraph_rtl(p_head)

    # 2. Board Pattern OMR Block

    if data.is_board_pattern:

        p_bhdr = doc.add_paragraph()

        p_bhdr.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_bhdr.add_run("FEDERAL BOARD / FGEI BOARD PATTERN EXAMINATION PAPER\n").bold = True

        table = doc.add_table(rows=2, cols=2)

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        table.cell(0, 0).text = "Roll No In Figures: __________________"

        table.cell(0, 1).text = "Sig. of Invigilator: __________________"

        table.cell(1, 0).text = "Roll No In Words: ___________________"

        table.cell(1, 1).text = "OMR Sheet No: ______________________"

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Divider

    p_hr = doc.add_paragraph()

    p_hr.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_hr.add_run("―" * 55).bold = True

    # 3. Section A: MCQs

    if data.mcqs:

        p_sec_a = doc.add_paragraph()

        sec_title = "حصہ اول (معروضی)" if data.is_urdu_paper else "Section – A (Multiple Choice Questions)"

        p_sec_a.add_run(sec_title).bold = True

        if data.is_urdu_paper:

            set_paragraph_rtl(p_sec_a)

        labels = ["(a)", "(b)", "(c)", "(d)"]

        for i, mcq in enumerate(data.mcqs, start=1):

            p_mcq = doc.add_paragraph()

            p_mcq.add_run(f"{i}.  ")

            add_text_with_math(p_mcq, mcq.text)

            p_mcq.add_run("\n  ")

            for opt_idx, opt_text in enumerate(mcq.options[:4]):

                p_mcq.add_run(f"{labels[opt_idx]} ")

                add_text_with_math(p_mcq, opt_text)

                p_mcq.add_run("   ")

            if data.is_urdu_paper or mcq.is_urdu:

                set_paragraph_rtl(p_mcq)

    # 4. Section B: Short Questions

    if data.short_questions:

        p_sec_b = doc.add_paragraph()

        sec_title = "حصہ دوم (مختصر سوالات)" if data.is_urdu_paper else "Section – B (Short Questions)"

        p_sec_b.add_run(sec_title).bold = True

        if data.is_urdu_paper:

            set_paragraph_rtl(p_sec_b)

        for i, sq in enumerate(data.short_questions, start=1):

            p_sq = doc.add_paragraph()

            p_sq.add_run(f"Q{i}. ").bold = True

            add_text_with_math(p_sq, sq.main_question)

            if sq.or_question:

                p_or = doc.add_paragraph()

                p_or.add_run("OR / یا: ").bold = True

                add_text_with_math(p_or, sq.or_question)

                if data.is_urdu_paper or sq.is_urdu:

                    set_paragraph_rtl(p_or)

            if data.is_urdu_paper or sq.is_urdu:

                set_paragraph_rtl(p_sq)

    # 5. Section C: Long Questions

    if data.long_questions:

        p_sec_c = doc.add_paragraph()

        sec_title = "حصہ سوم (تفصیلی سوالات)" if data.is_urdu_paper else "Section – C (Detailed / Long Questions)"

        p_sec_c.add_run(sec_title).bold = True

        if data.is_urdu_paper:

            set_paragraph_rtl(p_sec_c)

        offset = len(data.short_questions or [])

        for i, lq in enumerate(data.long_questions, start=1):

            p_lq = doc.add_paragraph()

            p_lq.add_run(f"Q{i+offset}. ").bold = True

            add_text_with_math(p_lq, lq.main_question)

            p_lq.add_run(f"  ({lq.marks} Marks)")

            if lq.or_question:

                p_or = doc.add_paragraph()

                p_or.add_run("OR / یا: ").bold = True

                add_text_with_math(p_or, lq.or_question)

                if data.is_urdu_paper or lq.is_urdu:

                    set_paragraph_rtl(p_or)

            if data.is_urdu_paper or lq.is_urdu:

                set_paragraph_rtl(p_lq)

    # 6. Section D: Project / Practical

    if data.project_questions:

        p_sec_d = doc.add_paragraph()

        sec_title = "عملی کام / پروجیکٹ" if data.is_urdu_paper else "Section – D (Practical / Project Section)"

        p_sec_d.add_run(sec_title).bold = True

        if data.is_urdu_paper:

            set_paragraph_rtl(p_sec_d)

        for i, pq in enumerate(data.project_questions, start=1):

            p_pq = doc.add_paragraph()

            p_pq.add_run(f"P{i}. ").bold = True

            add_text_with_math(p_pq, pq.main_question)

            p_pq.add_run(f"  ({pq.marks} Marks)")

            if data.is_urdu_paper or pq.is_urdu:

                set_paragraph_rtl(p_pq)

    doc.save(output_filename)

    return output_filename

def set_cell_text(cell, text, bold=False):
    cell.text = ""

    p = cell.paragraphs[0]

    run = p.add_run(text)
    run.bold = bold

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_table_borders(table):

    tbl = table._tbl

    tblPr = tbl.tblPr

    borders = OxmlElement('w:tblBorders')

    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):

        tag = 'w:' + edge

        element = OxmlElement(tag)

        element.set(
            qn('w:val'),
            'single'
        )

        element.set(
            qn('w:sz'),
            '4'
        )

        borders.append(element)

    tblPr.append(borders)

def set_cell_width(cell, width):
    cell.width = Inches(width)

def format_table_font(table, size=10):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(size)    

def get_question_number(section, index, previous_count=0):

    if section == "SHORT":
        return f"Q{index + previous_count}"

    elif section == "LONG":
        return f"Q{index + previous_count}"

    return f"Q{index}"

def create_ssc_short_table(doc, questions):

    table = doc.add_table(
        rows=1,
        cols=4
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
         "Q. No.",
         "Question",
         "OR",
         "Alternative Question"
    ]

    for i,h in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            h,
            True
        )

    for index,q in enumerate(questions,1):

        row = table.add_row().cells

        set_cell_text(
            row[0],
            str(index)
        )

        set_cell_text(
            row[1],
            q.main_question
        )

        if q.image:
            p = row[1].paragraphs[0]
            add_question_image(
                p,
                q.image
            )

        set_cell_text(
            row[2],
            "OR",
            True
        ) 

        set_cell_text(
            row[3],
            q.or_question or ""
        )

    widths=[
         0.5,
         3.0,
         0.5,
         3.0
     ]

    for row in table.rows:
        for i,w in enumerate(widths):
            set_cell_width(
                row.cells[i],
                w
            )  

    set_table_borders(table)
    format_table_font(table)

    return table

def create_ssc_long_table(doc, questions, start_number=3):
    table = doc.add_table(
        rows=1,
        cols=4
    )
    

    table.alignment = WD_TABLE_ALIGNMENT.CENTER


    headers = [
        "Q. No.",
        "Question",
        "OR",
        "Alternative Question"
    ]


    for i, h in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            h,
            True
        )


    for index, q in enumerate(questions, start_number):

        row = table.add_row().cells


        set_cell_text(
            row[0],
            f"Q{index}"
        )


        set_cell_text(
            row[1],
            q.main_question
        )

        if q.image:
            p = row[1].paragraphs[0]
            add_question_image(
                p,
                q.image
            )


        set_cell_text(
            row[2],
            "OR",
            True
        )


        set_cell_text(
            row[3],
            q.or_question or ""
        )


    widths = [
        0.6,
        3.0,
        0.5,
        3.0
    ]


    for row in table.rows:

        for i, width in enumerate(widths):

            set_cell_width(
                row.cells[i],
                width
            )


    set_table_borders(table)

    format_table_font(table)

    return table

def add_question_image(paragraph, image_path):

    if not image_path:
        return


    if os.path.exists(image_path):

        run = paragraph.add_run()

        run.add_picture(
            image_path,
            width=Inches(2.2)
        )  

def build_ssc_paper(data: ExamData, output_filename="fgei_paper.docx"):
    doc = Document()

    # Page Setup
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Logo
    logo_path = os.path.join("static", "logo.png")

    if os.path.exists(logo_path):

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(
            logo_path,
            width=Inches(0.8)
        )

    # Header

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(
        data.school_name
    )

    run.bold = True
    run.font.size = Pt(14)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p2.add_run(
        data.exam_title
    )

    run.bold = True
    run.font.size = Pt(12)

    # Information Table

    info = doc.add_table(
        rows=2,
        cols=4
    )

    info.alignment = WD_TABLE_ALIGNMENT.CENTER

    values = [
        "Subject",
        data.subject,
        "Class",
        data.grade,

        "Time",
        data.time_allowed,
        "Total Marks",
        str(data.total_marks)
    ]

    counter = 0

    for row in info.rows:
        for cell in row.cells:
            set_cell_text(
            cell,
            values[counter],
            counter % 2 == 0
        )
            counter += 1

    set_table_borders(info)        

    # Roll Number Area

    table = doc.add_table(
        rows=2,
        cols=2
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER


    table.cell(0,0).text = (
        "Roll No: __________________"
    )

    table.cell(0,1).text = (
        "Date: __________________"
    )


    table.cell(1,0).text = (
        "Signature: ______________"
    )

    table.cell(1,1).text = (
        "Invigilator: ____________"
    )

    doc.add_paragraph()


    # Divider

    p = doc.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p.add_run(
        "―"*55
    ).bold = True

    # Section A
    if data.mcqs:
        sec = doc.add_paragraph()

        sec.add_run(
            "Section A (Objective)"
        ).bold = True

        table = doc.add_table(
            rows=1,
            cols=7
        )

        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = [
            "No",
            "Question",
            "(A)",
            "(B)",
            "(C)",
            "(D)",
            "Response"
        ]

        for i,h in enumerate(headers):

            set_cell_text(
                table.rows[0].cells[i],
                h,
                True
            )
        for index, mcq in enumerate(data.mcqs,1):

            row = table.add_row().cells

            set_cell_text(
                row[0],
                str(index)
            )

            set_cell_text(
                row[1],
                mcq.text
            )

            for x in range(4):
                if x < len(mcq.options):
                    set_cell_text(
                        row[x+2],
                        mcq.options[x]
                    )

            set_cell_text(
                row[6],
                "○  ○  ○  ○"
            ) 

        set_table_borders(table)               
    
    # Section B
    if data.short_questions:
        sec = doc.add_paragraph()
        sec.alignment = WD_ALIGN_PARAGRAPH.CENTER

        total_short_marks = sum(
            q.marks or 0
            for q in data.short_questions
        )

        run = sec.add_run(
            f"Section B (Short Questions) "
            f"({total_short_marks} Marks)"
        )

        run.bold = True

        instruction = doc.add_paragraph()

        instruction.add_run(
            "Attempt the following questions."
        ).bold = True

        create_ssc_short_table(
            doc,
            data.short_questions
        )

    # Section C
    if data.long_questions:

        sec = doc.add_paragraph()
        sec.alignment = WD_ALIGN_PARAGRAPH.CENTER

        total_long_marks = sum(
                q.marks or 0
                for q in data.long_questions
            )

        run = sec.add_run(
                f"Section - C (Marks-{total_long_marks})"
            )
        run.bold = True

        note = doc.add_paragraph()

        note.add_run(
                "Note: Attempt all questions. "
                "Marks of all question are equal."
            ).bold = True

        start_number = (
                len(data.short_questions) + 2
            )

        create_ssc_long_table(
                doc,
                data.long_questions,
                start_number
            )
        
    doc.save(output_filename)

    return output_filename
@app.get("/")

def read_root():

    return FileResponse("static/index.html")

@app.post("/generate-paper")

def generate_paper(data: ExamData):

    try:

        output_file = build_fgei_paper(data)

        return FileResponse(

            path=output_file,

            filename=f"{data.subject}_{data.grade}_Exam.docx",

            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail=str(e)
        )